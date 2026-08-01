#!/usr/bin/env python3
"""
docs_export.py — convert generated feature-docs Markdown into PDF and/or DOCX.

Usage:
    python docs_export.py <markdown-file-or-glob> [<more files>...] [--formats pdf,docx] [--out-dir DIR]

Examples:
    python docs_export.py docs/specs/invoice-management-spec.md
    python docs_export.py "docs/specs/invoice-management-*.md" --formats pdf,docx
    python docs_export.py docs/specs/*.md --formats pdf

Requires: markdown, xhtml2pdf, python-docx
    pip install markdown xhtml2pdf python-docx
"""

import argparse
import glob
import os
import re
import sys

PDF_CSS = """
<style>
  body { font-family: Helvetica, Arial, sans-serif; font-size: 10.5pt; color: #222; }
  h1 { font-size: 20pt; border-bottom: 2px solid #F54E00; padding-bottom: 6px; }
  h2 { font-size: 15pt; color: #23251d; margin-top: 18px; border-left: 4px solid #F54E00; padding-left: 8px; }
  h3 { font-size: 12.5pt; color: #23251d; margin-top: 12px; }
  table { border-collapse: collapse; width: 100%; margin: 8px 0; }
  th, td { border: 1px solid #d4d2c9; padding: 4px 8px; font-size: 9pt; text-align: left; }
  th { background: #f4f3ee; }
  code { background: #f4f3ee; padding: 1px 4px; font-family: Courier, monospace; font-size: 9pt; }
  ul, ol { margin: 4px 0 8px 18px; }
  li { margin-bottom: 3px; }
</style>
"""


def strip_frontmatter(text: str) -> str:
    """Strip a leading YAML frontmatter block (--- ... ---), if present."""
    if text.startswith("---"):
        m = re.match(r"^---\s*\n.*?\n---\s*\n", text, re.DOTALL)
        if m:
            return text[m.end():].lstrip("\n")
    return text


def md_to_pdf(md_path: str, out_path: str) -> None:
    import markdown as md
    from xhtml2pdf import pisa

    with open(md_path, encoding="utf-8") as f:
        text = strip_frontmatter(f.read())

    body_html = md.markdown(text, extensions=["extra", "sane_lists", "tables"])
    full_html = f"<html><head><meta charset='utf-8'/>{PDF_CSS}</head><body>{body_html}</body></html>"

    with open(out_path, "wb") as out_f:
        result = pisa.CreatePDF(full_html, dest=out_f)
    if result.err:
        raise RuntimeError(f"xhtml2pdf reported {result.err} error(s) converting {md_path}")


def md_to_docx(md_path: str, out_path: str) -> None:
    from docx import Document
    from docx.shared import Pt

    with open(md_path, encoding="utf-8") as f:
        lines = strip_frontmatter(f.read()).splitlines()

    doc = Document()
    style = doc.styles["Normal"]
    style.font.size = Pt(10.5)

    table_buffer = []

    def flush_table():
        if not table_buffer:
            return
        rows = [r for r in table_buffer if not re.match(r"^\|?\s*:?-{2,}", r)]
        cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
        if cells:
            t = doc.add_table(rows=len(cells), cols=len(cells[0]))
            t.style = "Light Grid Accent 1"
            for ri, row in enumerate(cells):
                for ci, val in enumerate(row):
                    if ci < len(t.columns):
                        cell_p = t.cell(ri, ci).paragraphs[0]
                        add_inline(cell_p, val)
        table_buffer.clear()

    def add_inline(paragraph, text):
        # bold **x**, code `x`, plain text — good enough for generated docs
        tokens = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
        for tok in tokens:
            if not tok:
                continue
            if tok.startswith("**") and tok.endswith("**"):
                paragraph.add_run(tok[2:-2]).bold = True
            elif tok.startswith("`") and tok.endswith("`"):
                r = paragraph.add_run(tok[1:-1])
                r.font.name = "Courier New"
            else:
                paragraph.add_run(tok)

    for raw in lines:
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith("|"):
            table_buffer.append(stripped)
            continue
        else:
            flush_table()

        if not stripped:
            continue
        if stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith(("- ", "* ")):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, stripped[2:])
        elif re.match(r"^\d+\.\s", stripped):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\.\s+", "", stripped))
        elif stripped == "---":
            doc.add_paragraph("_" * 40)
        else:
            p = doc.add_paragraph()
            add_inline(p, stripped)

    flush_table()
    doc.save(out_path)


def main():
    ap = argparse.ArgumentParser(description="Convert feature-docs Markdown to PDF/DOCX")
    ap.add_argument("inputs", nargs="+", help="Markdown file(s) or glob pattern(s)")
    ap.add_argument("--formats", default="pdf", help="Comma-separated: pdf,docx (default: pdf)")
    ap.add_argument("--out-dir", default=None, help="Output directory (default: alongside source .md)")
    args = ap.parse_args()

    formats = [f.strip().lower() for f in args.formats.split(",") if f.strip()]
    unknown = set(formats) - {"pdf", "docx"}
    if unknown:
        print(f"  ERROR: unknown format(s): {', '.join(unknown)} (supported: pdf, docx)")
        sys.exit(1)

    files = []
    for pattern in args.inputs:
        matches = glob.glob(pattern)
        files.extend(matches if matches else ([pattern] if os.path.exists(pattern) else []))

    if not files:
        print("  ERROR: no matching Markdown files found.")
        sys.exit(1)

    for md_path in files:
        base = os.path.splitext(os.path.basename(md_path))[0]
        out_dir = args.out_dir or os.path.dirname(md_path) or "."
        os.makedirs(out_dir, exist_ok=True)

        for fmt in formats:
            out_path = os.path.join(out_dir, f"{base}.{fmt}")
            try:
                if fmt == "pdf":
                    md_to_pdf(md_path, out_path)
                elif fmt == "docx":
                    md_to_docx(md_path, out_path)
                print(f"  [OK] {md_path} -> {out_path}")
            except Exception as e:
                print(f"  [FAIL] {md_path} -> .{fmt}: {e}")


if __name__ == "__main__":
    main()
